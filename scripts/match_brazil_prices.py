#!/usr/bin/env python3
"""Match Brazil MSRP values from the prior global workbook to the refreshed Brazil data."""

from __future__ import annotations

import argparse
import json
import math
import re
import statistics
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from pathlib import Path

import openpyxl


NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
SOURCE_FIELDS = {
    "Country/Territory-Name",
    "Global Sales Sub-Segment",
    "Model (World)",
    "Model",
    "Energy Type",
    "Price. (local)",
    "Price. (EUR)",
}
YEARS = {2024: "CY2024", 2025: "CY2025", 2026: "CY2026"}
ALIASES = {
    ("dongfeng dfsk seres 3", "bev", "suv"): ("seres 3", "bev", "suv"),
}


def normalize(value: object) -> str:
    text = str(value or "").casefold().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return " ".join(text.split())


def as_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(",", "")
        if not text or text.casefold() in {"unspecified", "n/a", "na", "-"}:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
    return number if math.isfinite(number) else None


def cell_col(reference: str) -> int:
    letters = re.match(r"[A-Z]+", reference).group(0)
    value = 0
    for letter in letters:
        value = value * 26 + ord(letter) - 64
    return value - 1


def read_shared_strings(book: zipfile.ZipFile) -> list[str]:
    values: list[str] = []
    with book.open("xl/sharedStrings.xml") as stream:
        for _, element in ET.iterparse(stream, events=("end",)):
            if element.tag == f"{NS}si":
                values.append("".join(node.text or "" for node in element.iter(f"{NS}t")))
                element.clear()
    return values


def value_of(cell: ET.Element, shared: list[str]):
    kind = cell.attrib.get("t")
    value = cell.find(f"{NS}v")
    if value is None or value.text is None:
        inline = cell.find(f"{NS}is")
        return "" if inline is None else "".join(node.text or "" for node in inline.iter(f"{NS}t"))
    raw = value.text
    if kind == "s":
        return shared[int(raw)]
    if kind in {"str", "inlineStr"}:
        return raw
    try:
        number = float(raw)
        return int(number) if number.is_integer() else number
    except ValueError:
        return raw


def read_prior_prices(source: Path) -> tuple[dict[tuple[str, str, str], list[float]], dict[tuple[str, str], list[float]], list[float]]:
    exact: dict[tuple[str, str, str], list[float]] = defaultdict(list)
    model_fuel: dict[tuple[str, str], list[float]] = defaultdict(list)
    fx_rates: list[float] = []
    with zipfile.ZipFile(source) as book:
        shared = read_shared_strings(book)
        needed_cols: dict[int, str] = {}
        with book.open("xl/worksheets/sheet7.xml") as stream:
            for _, element in ET.iterparse(stream, events=("end",)):
                if element.tag != f"{NS}row":
                    continue
                row_number = int(element.attrib.get("r", "0"))
                cells = {cell_col(cell.attrib["r"]): value_of(cell, shared) for cell in element.findall(f"{NS}c")}
                if row_number == 1:
                    headers = {index: str(value).strip() for index, value in cells.items()}
                    needed_cols = {index: name for index, name in headers.items() if name in SOURCE_FIELDS}
                    missing = SOURCE_FIELDS.difference(needed_cols.values())
                    if missing:
                        raise RuntimeError(f"Missing prior-source columns: {sorted(missing)}")
                    element.clear()
                    continue
                row = {name: cells.get(index, "") for index, name in needed_cols.items()}
                if str(row["Country/Territory-Name"]).strip() != "Brazil":
                    element.clear()
                    continue
                model = normalize(row["Model (World)"])
                fuel = normalize(row["Energy Type"])
                segment = normalize(row["Global Sales Sub-Segment"])
                local_price = as_number(row["Price. (local)"])
                eur_price = as_number(row["Price. (EUR)"])
                if model and fuel and segment and local_price and local_price > 0:
                    exact[(model, fuel, segment)].append(local_price)
                    model_fuel[(model, fuel)].append(local_price)
                    if eur_price and eur_price > 0:
                        fx_rates.append(local_price / eur_price)
                element.clear()
    return exact, model_fuel, fx_rates


def unique_median(values: list[float]) -> int:
    return round(statistics.median(sorted(set(values))))


def safe_id(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.casefold()).strip("-")


def brand_of(model: str) -> str:
    two_word = ("Alfa Romeo", "Aston Martin", "Land Rover")
    for brand in two_word:
        if model.startswith(brand + " "):
            return brand
    return model.split()[0] if model.split() else ""


def load_manual_prices(path: Path | None) -> dict[tuple[str, str, str], dict]:
    if not path:
        return {}
    rows = json.loads(path.read_text(encoding="utf-8"))
    result = {}
    for row in rows:
        prices = [number for value in row.get("prices", []) if (number := as_number(value)) and number > 0]
        if not prices:
            continue
        key = (normalize(row["model"]), normalize(row["fuel"]), normalize(row["segment"]))
        result[key] = {
            "price": unique_median(prices),
            "priceMin": round(min(prices)),
            "priceMax": round(max(prices)),
            "source": "网络补充",
            "sourceUrl": row.get("sourceUrl", ""),
            "note": row.get("note", ""),
        }
    return result


def resolve_price(
    key: tuple[str, str, str],
    current_price: float | None,
    exact: dict[tuple[str, str, str], list[float]],
    model_fuel: dict[tuple[str, str], list[float]],
    manual: dict[tuple[str, str, str], dict],
) -> dict | None:
    if current_price and current_price > 0:
        price = round(current_price)
        return {"price": price, "priceMin": price, "priceMax": price, "source": "新表原值", "sourceUrl": "", "note": "新表已有价格"}
    if key in exact:
        values = exact[key]
        return {"price": unique_median(values), "priceMin": round(min(values)), "priceMax": round(max(values)), "source": "旧版合并表-车型/动力/车身精确匹配", "sourceUrl": "", "note": "多配置取中位数"}
    alias = ALIASES.get(key)
    if alias and alias in exact:
        values = exact[alias]
        return {"price": unique_median(values), "priceMin": round(min(values)), "priceMax": round(max(values)), "source": "旧版合并表-别名匹配", "sourceUrl": "", "note": f"别名匹配：{' / '.join(alias)}；多配置取中位数"}
    model, fuel, _segment = key
    if (model, fuel) in model_fuel:
        values = model_fuel[(model, fuel)]
        return {"price": unique_median(values), "priceMin": round(min(values)), "priceMax": round(max(values)), "source": "旧版合并表-车型/动力匹配", "sourceUrl": "", "note": "跨车身记录匹配；多配置取中位数"}
    return manual.get(key)


def build_outputs(current: Path, prior: Path, manual_path: Path | None) -> tuple[dict, dict, list[dict], float | None]:
    exact, model_fuel, fx_rates = read_prior_prices(prior)
    manual = load_manual_prices(manual_path)
    workbook = openpyxl.load_workbook(current, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    required = {"Country/Territory-Name", "Global Sales Sub-Segment", "Model (World)", "Fuel Type", "Price", "CY2024", "CY2025", "CY2026"}
    missing_headers = required.difference(headers)
    if missing_headers:
        raise RuntimeError(f"Missing current columns: {sorted(missing_headers)}")
    index = {name: position for position, name in enumerate(headers)}
    keys: dict[tuple[str, str, str], dict] = {}
    patch_rows: list[list[object]] = []
    site_groups: dict[tuple[int, str, str, str], dict] = defaultdict(
        lambda: {"sales": 0.0, "lengths": [], "prices": [], "priceMins": [], "priceMaxs": [], "sourceRows": 0}
    )
    positive_rows = 0
    existing_price_rows = 0
    for values in rows:
        country = str(values[index["Country/Territory-Name"]] or "").strip()
        segment_raw = str(values[index["Global Sales Sub-Segment"]] or "").strip()
        model_raw = str(values[index["Model (World)"]] or "").strip()
        fuel_raw = str(values[index["Fuel Type"]] or "").strip()
        current_price = as_number(values[index["Price"]])
        if country != "Brazil" or segment_raw not in {"Car", "SUV", "MPV"}:
            patch_rows.append([current_price, "", "", ""])
            continue
        if not model_raw or model_raw == "Unspecified" or not fuel_raw or fuel_raw == "Unspecified":
            note = "未参与匹配：车型或动力形式为 Unspecified"
            patch_rows.append([current_price, "未匹配", "", note])
            continue
        key = (normalize(model_raw), normalize(fuel_raw), normalize(segment_raw))
        resolved = resolve_price(key, current_price, exact, model_fuel, manual)
        if resolved:
            patch_rows.append([resolved["price"], resolved["source"], resolved["sourceUrl"], resolved["note"]])
        else:
            patch_rows.append([None, "未找到巴西公开 MSRP", "", "已检索巴西官方渠道及本地汽车媒体；不使用海外售价或非官方进口报价"])
        sales = sum(as_number(values[index[year]]) or 0 for year in ("CY2024", "CY2025", "CY2026"))
        if sales <= 0:
            continue
        positive_rows += 1
        if current_price and current_price > 0:
            existing_price_rows += 1
        entry = keys.setdefault(key, {"model": model_raw, "fuel": fuel_raw, "segment": segment_raw, "sales": 0.0, "rows": 0, "existingPrices": [], "lengths": set(), "generationYears": set(), "salesByYear": defaultdict(float)})
        entry["sales"] += sales
        entry["rows"] += 1
        length = as_number(values[index.get("Length")]) if "Length" in index else None
        generation_year = as_number(values[index.get("Generation Year")]) if "Generation Year" in index else None
        if length and length > 0:
            entry["lengths"].add(round(length))
        if generation_year and generation_year > 0:
            entry["generationYears"].add(round(generation_year))
        for year in ("CY2024", "CY2025", "CY2026"):
            entry["salesByYear"][year] += as_number(values[index[year]]) or 0
        if current_price and current_price > 0:
            entry["existingPrices"].append(current_price)
        if resolved:
            length = as_number(values[index["Length"]]) if "Length" in index else None
            for year, header in YEARS.items():
                year_sales = as_number(values[index[header]]) or 0
                if year_sales <= 0:
                    continue
                group = site_groups[(year, model_raw, fuel_raw, segment_raw)]
                group["sales"] += year_sales
                if length and length > 0:
                    group["lengths"].append(length)
                group["prices"].append(resolved["price"])
                group["priceMins"].append(resolved["priceMin"])
                group["priceMaxs"].append(resolved["priceMax"])
                group["sourceRows"] += 1
    workbook.close()

    matched = []
    missing = []
    for key, entry in sorted(keys.items(), key=lambda item: (-item[1]["sales"], item[1]["model"])):
        entry["lengths"] = sorted(entry["lengths"])
        entry["generationYears"] = sorted(entry["generationYears"])
        entry["salesByYear"] = {year: round(value) for year, value in entry["salesByYear"].items() if value > 0}
        resolved = resolve_price(key, unique_median(entry["existingPrices"]) if entry["existingPrices"] else None, exact, model_fuel, manual)
        if resolved:
            entry.update(resolved)
            matched.append(entry)
        else:
            missing.append(entry)

    site_records = []
    for (year, model, fuel, segment), group in sorted(site_groups.items()):
        sensible_lengths = [value for value in group["lengths"] if value >= 3000]
        lengths = sensible_lengths or group["lengths"]
        if not lengths:
            continue
        site_records.append({
            "id": f"br-{year}-{safe_id(model)}-{safe_id(fuel)}-{safe_id(segment)}",
            "country": "Brazil",
            "year": year,
            "brand": brand_of(model),
            "model": model,
            "localModels": [],
            "fuel": fuel,
            "body": segment,
            "length": round(statistics.median(lengths)),
            "price": round(statistics.median(group["prices"])),
            "priceMin": round(min(group["priceMins"])),
            "priceMax": round(max(group["priceMaxs"])),
            "sales": round(group["sales"]),
            "sourceRows": group["sourceRows"],
        })

    report = {
        "currentWorkbook": current.name,
        "priorWorkbook": prior.name,
        "currentRowsWithSales": positive_rows,
        "currentRowsAlreadyPriced": existing_price_rows,
        "uniqueVehicleKeys": len(keys),
        "matchedVehicleKeys": len(matched),
        "missingVehicleKeys": len(missing),
        "medianLocalPerEuro": round(statistics.median(fx_rates), 6) if fx_rates else None,
        "matchedBySource": {source: sum(1 for row in matched if row["source"] == source) for source in sorted({row["source"] for row in matched})},
        "siteRows": len(site_records),
        "missing": missing,
        "matched": matched,
    }
    patch = {
        "sheetName": sheet.title,
        "sourceRowCount": sheet.max_row,
        "headers": ["Price", "价格来源", "来源链接", "价格备注"],
        "rows": patch_rows,
    }
    return report, patch, site_records, report["medianLocalPerEuro"]


def write_site_data(path: Path, records: list[dict], report: dict, fx_rate: float | None):
    fx = fx_rate or 6.388832
    references = [
        {"id": "reference-s36a", "model": "S36A", "length": 4500, "price": round(30000 * fx), "reference": True},
        {"id": "reference-s36", "model": "S36", "length": 4650, "price": round(35000 * fx), "reference": True},
    ]
    meta = {
        "country": "Brazil",
        "years": list(YEARS),
        "latestYear": 2026,
        "latestPeriod": "2026 YTD (Jan–Jul)",
        "rows": len(records),
        "sourceWorkbook": report["currentWorkbook"],
        "sourceSheet": "Sheet1",
        "currency": "BRL",
        "referenceFxBrlPerEur": fx,
        "method": "Car/SUV/MPV only; Unspecified removed; Brazil MSRP matched from prior workbook then supplemented by web research; multiple configurations use median; sales summed by year + Model (World) + fuel + segment.",
        "priceCoverage": {"matchedVehicleKeys": report["matchedVehicleKeys"], "missingVehicleKeys": report["missingVehicleKeys"]},
    }
    payload = (
        "window.VEHICLE_DATA = " + json.dumps(records, ensure_ascii=False, separators=(",", ":"))
        + ";\nwindow.REFERENCE_DATA = " + json.dumps(references, ensure_ascii=False, separators=(",", ":"))
        + ";\nwindow.DATA_META = " + json.dumps(meta, ensure_ascii=False, separators=(",", ":")) + ";\n"
    )
    path.write_text(payload, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("current", type=Path)
    parser.add_argument("prior", type=Path)
    parser.add_argument("--manual", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--patch", type=Path)
    parser.add_argument("--site", type=Path)
    args = parser.parse_args()
    report, patch, site_records, fx_rate = build_outputs(args.current, args.prior, args.manual)
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        args.report.write_text(payload, encoding="utf-8")
    if args.patch:
        args.patch.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")
    if args.site:
        write_site_data(args.site, site_records, report, fx_rate)
    print(payload)


if __name__ == "__main__":
    main()
